"""文档生成：把 AI 回答的 Markdown 转成可下载的文件。

支持格式：Word (.docx) / PDF / HTML / Markdown / 纯文本
"""

from __future__ import annotations

import html
import io
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "exports"
OUTPUT_DIR.mkdir(exist_ok=True)

FENCE = chr(96) * 3
NL = chr(10)


@dataclass
class Block:
    kind: str
    text: str = ""
    lang: str = ""
    level: int = 0
    ordered: bool = False
    rows: list = field(default_factory=list)


def parse_markdown(md: str) -> list:
    """把 Markdown 解析成块列表。"""
    blocks = []
    lines = (md or "").replace(chr(13) + chr(10), NL).split(NL)
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith(FENCE):
            lang = stripped[3:].strip()
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith(FENCE):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", NL.join(code_lines), lang=lang))
            continue

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.match(r"^[-*_]{3,}$", stripped):
            blocks.append(Block("hr"))
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})[ ]+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            blocks.append(Block("h" + str(min(level, 3)), m.group(2).strip(), level=level))
            i += 1
            continue

        # 表格
        if "|" in line and i + 1 < n and re.match(r"^[ ]*\|?[\s:|-]+\|[\s:|-]*$", lines[i + 1]):
            def split_row(row):
                r = row.strip()
                if r.startswith("|"):
                    r = r[1:]
                if r.endswith("|"):
                    r = r[:-1]
                return [c.strip() for c in r.split("|")]

            header = split_row(line)
            i += 2
            rows = [header]
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(Block("table", rows=rows))
            continue

        # 引用
        if stripped.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(Block("quote", " ".join(quote)))
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+\.)[ ]+(.*)$", line)
        if m:
            indent = len(m.group(1))
            ordered = m.group(2)[0].isdigit()
            blocks.append(Block("li", m.group(3).strip(), level=indent // 2, ordered=ordered))
            i += 1
            continue

        # 段落
        para = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if not nxt:
                break
            if re.match(r"^(#{1,6}[ ]|[-*+][ ]|\d+\.[ ]|>|\|)", nxt):
                break
            if nxt.startswith(FENCE):
                break
            para.append(nxt)
            i += 1
        blocks.append(Block("p", " ".join(para)))

    return blocks


def _clean_inline(text: str) -> str:
    """去掉行内 Markdown 标记。"""
    text = re.sub(r"\*\*([^*]+)\*\*", lambda m: m.group(1), text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", lambda m: m.group(1), text)
    text = re.sub(r"[`]([^`]+)[`]", lambda m: m.group(1), text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)",
                  lambda m: m.group(1) + " (" + m.group(2) + ")", text)
    return text.strip()


# ============ Markdown / 纯文本 ============

def to_markdown(title, body, meta=None):
    meta = meta or {}
    lines = ["# " + title, ""]
    for k, v in meta.items():
        lines.append("- **" + str(k) + "**：" + str(v))
    if meta:
        lines.extend(["", "---", ""])
    lines.append(body)
    lines.extend(["", "---", "*导出于 " + time.strftime("%Y-%m-%d %H:%M:%S") + "*"])
    return NL.join(lines)


def to_text(title, body):
    out = [title, "=" * 40, ""]
    for b in parse_markdown(body):
        if b.kind.startswith("h"):
            t = _clean_inline(b.text)
            out.extend(["", t, "-" * max(len(t), 3)])
        elif b.kind == "code":
            out.append(b.text)
        elif b.kind == "li":
            prefix = "- " if not b.ordered else "• "
            out.append("  " * b.level + prefix + _clean_inline(b.text))
        elif b.kind == "table":
            for row in b.rows:
                out.append(" | ".join(_clean_inline(c) for c in row))
        elif b.kind == "quote":
            out.append("> " + _clean_inline(b.text))
        elif b.kind == "hr":
            out.append("")
        else:
            out.append(_clean_inline(b.text))
        out.append("")
    return NL.join(out)


# ============ HTML ============

HTML_CSS = """
  :root { --fg:#18181b; --dim:#71717a; --bd:#e4e4e7; --accent:#5145e5; }
  * { box-sizing: border-box; }
  body {
    font-family: -apple-system, "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif;
    max-width: 820px; margin: 0 auto; padding: 48px 28px 80px;
    color: var(--fg); line-height: 1.75; font-size: 15.5px; background: #fff;
  }
  .doc-title { font-size: 30px; font-weight: 700; margin: 0 0 8px; letter-spacing: -.02em; }
  .doc-sub { color: var(--dim); font-size: 13px; margin-bottom: 24px; }
  .meta {
    display: flex; flex-wrap: wrap; gap: 20px;
    padding: 14px 18px; background: #fafafa;
    border: 1px solid var(--bd); border-radius: 10px; margin-bottom: 32px;
  }
  .meta-item { font-size: 12.5px; color: var(--dim); }
  .meta-item span { margin-right: 6px; }
  .meta-item strong { color: var(--fg); font-weight: 600; }
  h1, h2, h3 { margin: 32px 0 14px; line-height: 1.35; letter-spacing: -.01em; }
  h1 { font-size: 25px; } h2 { font-size: 20px; } h3 { font-size: 17px; }
  p { margin: 0 0 14px; }
  ul, ol { margin: 0 0 14px; padding-left: 26px; }
  li { margin: 5px 0; }
  blockquote {
    margin: 0 0 16px; padding: 10px 18px;
    border-left: 3px solid var(--accent); background: #fafafa;
    color: #52525b; border-radius: 0 6px 6px 0;
  }
  code {
    background: #f4f4f5; padding: 2px 6px; border-radius: 5px;
    font-family: "SF Mono", Consolas, monospace; font-size: .88em; color: #4c3fd4;
  }
  pre {
    background: #1c1c22; color: #e4e4e9; padding: 18px 20px;
    border-radius: 10px; overflow-x: auto; margin: 0 0 18px;
    font-size: 13.5px; line-height: 1.6;
  }
  pre code { background: none; color: inherit; padding: 0; font-size: inherit; }
  table { border-collapse: collapse; width: 100%; margin: 0 0 18px; font-size: 14px; }
  th, td { border: 1px solid var(--bd); padding: 9px 13px; text-align: left; }
  th { background: #fafafa; font-weight: 600; }
  hr { border: none; border-top: 1px solid var(--bd); margin: 28px 0; }
  .footer {
    margin-top: 48px; padding-top: 18px; border-top: 1px solid var(--bd);
    color: var(--dim); font-size: 12px; text-align: center;
  }
  @media print { body { padding: 0; max-width: none; } }
"""


def to_html(title, body, meta=None):
    meta = meta or {}
    parts = []
    state = {"in_list": False, "ordered": False}

    def close_list():
        if state["in_list"]:
            parts.append("</ol>" if state["ordered"] else "</ul>")
            state["in_list"] = False

    for b in parse_markdown(body):
        if b.kind != "li":
            close_list()

        if b.kind in ("h1", "h2", "h3"):
            parts.append("<" + b.kind + ">" + html.escape(b.text) + "</" + b.kind + ">")
        elif b.kind == "code":
            parts.append("<pre><code>" + html.escape(b.text) + "</code></pre>")
        elif b.kind == "quote":
            parts.append("<blockquote>" + html.escape(b.text) + "</blockquote>")
        elif b.kind == "hr":
            parts.append("<hr>")
        elif b.kind == "table":
            t = ["<table><thead><tr>"]
            for c in b.rows[0]:
                t.append("<th>" + html.escape(_clean_inline(c)) + "</th>")
            t.append("</tr></thead><tbody>")
            for row in b.rows[1:]:
                t.append("<tr>")
                for c in row:
                    t.append("<td>" + html.escape(_clean_inline(c)) + "</td>")
                t.append("</tr>")
            t.append("</tbody></table>")
            parts.append("".join(t))
        elif b.kind == "li":
            if not state["in_list"]:
                state["ordered"] = b.ordered
                parts.append("<ol>" if b.ordered else "<ul>")
                state["in_list"] = True
            parts.append("<li>" + html.escape(_clean_inline(b.text)) + "</li>")
        else:
            parts.append("<p>" + html.escape(_clean_inline(b.text)) + "</p>")

    close_list()

    meta_html = ""
    if meta:
        items = "".join(
            '<div class="meta-item"><span>' + html.escape(str(k)) + "</span><strong>"
            + html.escape(str(v)) + "</strong></div>"
            for k, v in meta.items()
        )
        meta_html = '<div class="meta">' + items + "</div>"

    return (
        '<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        "<title>" + html.escape(title) + "</title><style>" + HTML_CSS + "</style></head><body>"
        '<div class="doc-title">' + html.escape(title) + "</div>"
        '<div class="doc-sub">导出于 ' + time.strftime("%Y-%m-%d %H:%M:%S") + "</div>"
        + meta_html + "".join(parts)
        + '<div class="footer">由 AI Agent 演示生成</div></body></html>'
    )


# ============ Word ============

def to_docx(title, body, meta=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    meta = meta or {}
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    if meta:
        p = doc.add_paragraph()
        for k, v in meta.items():
            r1 = p.add_run(str(k) + "：")
            r1.font.size = Pt(9)
            r1.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)
            r2 = p.add_run(str(v) + "    ")
            r2.font.size = Pt(9)
        doc.add_paragraph()

    for b in parse_markdown(body):
        if b.kind in ("h1", "h2", "h3"):
            doc.add_heading(_clean_inline(b.text), level=int(b.kind[1]))
        elif b.kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1C, 0x1C, 0x22)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F4F4F5")
            p._p.get_or_add_pPr().append(shd)
        elif b.kind == "li":
            name = "List Number" if b.ordered else "List Bullet"
            try:
                doc.add_paragraph(_clean_inline(b.text), style=name)
            except KeyError:
                doc.add_paragraph("• " + _clean_inline(b.text))
        elif b.kind == "table" and b.rows:
            rows, cols = len(b.rows), len(b.rows[0])
            table = doc.add_table(rows=rows, cols=cols)
            try:
                table.style = "Light Grid Accent 1"
            except KeyError:
                pass
            for ri, row in enumerate(b.rows):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    cell.text = _clean_inline(row[ci]) if ci < len(row) else ""
                    if ri == 0:
                        for p2 in cell.paragraphs:
                            for r2 in p2.runs:
                                r2.font.bold = True
        elif b.kind == "quote":
            p = doc.add_paragraph(_clean_inline(b.text))
            p.paragraph_format.left_indent = Inches(0.3)
            for r in p.runs:
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x52, 0x52, 0x5B)
        elif b.kind == "hr":
            doc.add_paragraph("—" * 20)
        else:
            doc.add_paragraph(_clean_inline(b.text))

    doc.add_paragraph()
    foot = doc.add_paragraph("导出于 " + time.strftime("%Y-%m-%d %H:%M:%S"))
    for r in foot.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x9C, 0x9C, 0xA5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============ PDF ============

def to_pdf(title, body):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer

    font_name = "Helvetica"
    for cand in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc",
                 "C:/Windows/Fonts/simhei.ttf"]:
        if Path(cand).exists():
            try:
                pdfmetrics.registerFont(TTFont("CJK", cand))
                font_name = "CJK"
                break
            except Exception:
                continue

    base = ParagraphStyle("b", fontName=font_name, fontSize=10.5, leading=17, spaceAfter=6)
    h1s = ParagraphStyle("h1s", parent=base, fontSize=18, leading=26,
                         spaceBefore=14, spaceAfter=8)
    h2s = ParagraphStyle("h2s", parent=base, fontSize=14.5, leading=22,
                         spaceBefore=12, spaceAfter=6)
    h3s = ParagraphStyle("h3s", parent=base, fontSize=12, leading=19,
                         spaceBefore=10, spaceAfter=5)
    code_s = ParagraphStyle("c", parent=base, fontSize=8.5, leading=13,
                            leftIndent=10, backColor=colors.HexColor("#F4F4F5"),
                            borderPadding=6, spaceBefore=4, spaceAfter=8)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm, title=title)
    story = [Paragraph(html.escape(title), h1s), Spacer(1, 6)]

    for b in parse_markdown(body):
        if b.kind == "h1":
            story.append(Paragraph(html.escape(_clean_inline(b.text)), h1s))
        elif b.kind == "h2":
            story.append(Paragraph(html.escape(_clean_inline(b.text)), h2s))
        elif b.kind == "h3":
            story.append(Paragraph(html.escape(_clean_inline(b.text)), h3s))
        elif b.kind == "code":
            story.append(Preformatted(b.text, code_s))
        elif b.kind == "li":
            prefix = "· " if b.ordered else "• "
            story.append(Paragraph(prefix + html.escape(_clean_inline(b.text)), base))
        elif b.kind == "quote":
            qs = ParagraphStyle("q", parent=base, leftIndent=14,
                                textColor=colors.HexColor("#52525B"))
            story.append(Paragraph("<i>" + html.escape(_clean_inline(b.text)) + "</i>", qs))
        elif b.kind == "hr":
            story.append(Spacer(1, 10))
        elif b.kind == "table":
            for row in b.rows:
                story.append(Paragraph(
                    " | ".join(html.escape(_clean_inline(c)) for c in row), base))
        else:
            story.append(Paragraph(html.escape(_clean_inline(b.text)), base))

    story.append(Spacer(1, 16))
    story.append(Paragraph(
        '<font size="8" color="#9C9CA5">导出于 '
        + time.strftime("%Y-%m-%d %H:%M:%S") + "</font>", base))
    doc.build(story)
    return buf.getvalue()


# ============ 统一入口 ============

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "txt": ("text/plain; charset=utf-8", "txt"),
    "html": ("text/html; charset=utf-8", "html"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pdf": ("application/pdf", "pdf"),
}


def build(fmt, title, body, meta=None):
    fmt = (fmt or "md").lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError("不支持的格式：" + fmt + "，可选：" + ", ".join(FORMATS))
    mime, ext = FORMATS[fmt]
    if fmt == "md":
        return to_markdown(title, body, meta).encode("utf-8"), mime, ext
    if fmt == "txt":
        return to_text(title, body).encode("utf-8"), mime, ext
    if fmt == "html":
        return to_html(title, body, meta).encode("utf-8"), mime, ext
    if fmt == "docx":
        return to_docx(title, body, meta), mime, ext
    return to_pdf(title, body), mime, ext


def safe_filename(name, ext):
    name = re.sub(r'[\\/:*?"<>|\r\n]+', "_", (name or "document")).strip()
    if not name:
        name = "document"
    if len(name) > 60:
        name = name[:60]
    return name + "_" + time.strftime("%Y%m%d_%H%M%S") + "." + ext
